#!/usr/bin/env python3

"""
 * Copyright 2026 NXP
 * All rights reserved.
 *
 * SPDX-License-Identifier: BSD-3-Clause
"""

"""
Coverity Issues Parser for MPP Project

This script parses Coverity Excel files and generates statistics based on:
1. Impact-based filtering (Medium/High)
2. Security classification filtering (Highest/High)
"""

import argparse
import pandas as pd
from pathlib import Path
from typing import List, Dict, Set
import sys
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font


class CoverityParser:
    def __init__(self, excel_files: List[str], debug: bool = False, show_intermediate: bool = False,
                 coverity_server: str = "coverity9.nxp.com", project: str = "MPP",
                 fail_on_security: bool = False, fail_on_triage: bool = False):
        self.excel_files = excel_files
        self.debug = debug
        self.show_intermediate = show_intermediate
        self.coverity_server = coverity_server
        self.project = project
        self.fail_on_security = fail_on_security
        self.fail_on_triage = fail_on_triage
        self.impact_stats: Dict[str, Set[int]] = {"Medium": set(), "High": set()}
        self.security_stats: Dict[str, Set[int]] = {"Highest": set(), "High": set()}
        self.triage_stats: Dict[str, Set[int]] = {"Triaged": set(), "Dismissed": set()}
        self.impact_issues: List[pd.Series] = []
        self.security_issues: List[pd.Series] = []
        self.triage_issues: List[pd.Series] = []

    def log_debug(self, message: str):
        """Print debug messages if debug mode is enabled"""
        if self.debug:
            print(f"[DEBUG] {message}")

    def log_info(self, message: str):
        """Print info messages"""
        print(f"[INFO] {message}")

    def log_error(self, message: str):
        """Print error messages"""
        print(f"[ERROR] {message}", file=sys.stderr)

    def read_excel(self, file_path: str) -> pd.DataFrame:
        """Read the Issues_Plain sheet from an Excel file"""
        self.log_info(f"Reading file: {file_path}")
        try:
            df = pd.read_excel(file_path, sheet_name="Issues_Plain")
            self.log_debug(f"Total rows read: {len(df)}")
            return df
        except Exception as e:
            print(f"[ERROR] Failed to read {file_path}: {e}")
            sys.exit(1)

    def filter_impact_based(self, df: pd.DataFrame) -> pd.DataFrame:
        """Apply impact-based filters in order"""
        self.log_info("Applying impact-based filters...")

        # Filter 1: displayFile contains 'mpp'
        df_filtered = df[df['displayFile'].str.contains('mpp', case=False, na=False)]
        self.log_debug(f"After displayFile filter (contains 'mpp'): {len(df_filtered)} rows")

        # Filter 2: displayImpact is Medium or High
        df_filtered = df_filtered[df_filtered['displayImpact'].isin(['Medium', 'High'])]
        self.log_debug(f"After displayImpact filter (Medium/High): {len(df_filtered)} rows")

        # Filter 3: Exclude status = Dismissed
        df_filtered = df_filtered[df_filtered['status'] != 'Dismissed']
        self.log_debug(f"After status filter (exclude Dismissed): {len(df_filtered)} rows")

        return df_filtered

    def filter_security_based(self, df: pd.DataFrame) -> pd.DataFrame:
        """Apply security classification-based filters in order"""
        self.log_info("Applying security classification-based filters...")

        # Filter 1: displayFile contains 'mpp'
        df_filtered = df[df['displayFile'].str.contains('mpp', case=False, na=False)]
        self.log_debug(f"After displayFile filter (contains 'mpp'): {len(df_filtered)} rows")

        # Filter 2: security_classification is 1 - Highest or 2 - High
        df_filtered = df_filtered[df_filtered['security_classification'].isin(['1 - Highest', '2 - High'])]
        self.log_debug(f"After security_classification filter (Highest/High): {len(df_filtered)} rows")

        # Filter 3: Exclude status = Dismissed
        df_filtered = df_filtered[df_filtered['status'] != 'Dismissed']
        self.log_debug(f"After status filter (exclude Dismissed): {len(df_filtered)} rows")

        return df_filtered

    def filter_triage_based(self, df: pd.DataFrame) -> pd.DataFrame:
        """Apply triage-based filters in order"""
        self.log_info("Applying triage-based filters...")

        # Filter 1: displayFile contains 'mpp'
        df_filtered = df[df['displayFile'].str.contains('mpp', case=False, na=False)]
        self.log_debug(f"After displayFile filter (contains 'mpp'): {len(df_filtered)} rows")

        # Filter 2: status is Triaged or Dismissed
        df_filtered = df_filtered[df_filtered['status'].isin(['Triaged', 'Dismissed'])]
        self.log_debug(f"After status filter (Triaged/Dismissed): {len(df_filtered)} rows")

        # Filter 3: lastTriageComment is empty (null or empty string)
        df_filtered = df_filtered[df_filtered['lastTriageComment'].isna() | (df_filtered['lastTriageComment'] == '')]
        self.log_debug(f"After lastTriageComment filter (empty): {len(df_filtered)} rows")

        return df_filtered

    def process_impact_stats(self, df: pd.DataFrame, file_name: str):
        """Process impact-based statistics"""
        self.log_info(f"Processing impact statistics for {file_name}...")

        for _, row in df.iterrows():
            impact = row['displayImpact']
            cid = int(row['cid'])

            if self.debug:
                self.log_debug(f"Impact issue found - CID: {cid}, Impact: {impact}")
                self.log_debug(f"Full row data:\n{row.to_dict()}")

            if impact in self.impact_stats:
                if cid not in self.impact_stats[impact]:
                    self.impact_stats[impact].add(cid)
                    self.impact_issues.append(row)

        if self.show_intermediate:
            self.print_intermediate_impact_stats(file_name)

    def process_security_stats(self, df: pd.DataFrame, file_name: str):
        """Process security classification-based statistics"""
        self.log_info(f"Processing security statistics for {file_name}...")

        for _, row in df.iterrows():
            security = row['security_classification']
            cid = int(row['cid'])

            if self.debug:
                self.log_debug(f"Security issue found - CID: {cid}, Classification: {security}")
                self.log_debug(f"Full row data:\n{row.to_dict()}")

            if security == '1 - Highest':
                if cid not in self.security_stats['Highest']:
                    self.security_stats['Highest'].add(cid)
                    self.security_issues.append(row)
            elif security == '2 - High':
                if cid not in self.security_stats['High']:
                    self.security_stats['High'].add(cid)
                    self.security_issues.append(row)

        if self.show_intermediate:
            self.print_intermediate_security_stats(file_name)

    def process_triage_stats(self, df: pd.DataFrame, file_name: str):
        """Process triage-based statistics"""
        self.log_info(f"Processing triage statistics for {file_name}...")

        for _, row in df.iterrows():
            status = row['status']
            cid = int(row['cid'])

            if self.debug:
                self.log_debug(f"Triage issue found - CID: {cid}, Status: {status}")
                self.log_debug(f"Full row data:\n{row.to_dict()}")

            if status in self.triage_stats:
                if cid not in self.triage_stats[status]:
                    self.triage_stats[status].add(cid)
                    self.triage_issues.append(row)

        if self.show_intermediate:
            self.print_intermediate_triage_stats(file_name)

    def print_intermediate_impact_stats(self, file_name: str):
        """Print intermediate impact statistics"""
        print(f"\n{'='*60}")
        print(f"Intermediate Impact Statistics for: {file_name}")
        print(f"{'='*60}")
        print(f"Medium Issues: {len(self.impact_stats['Medium'])}")
        print(f"High Issues: {len(self.impact_stats['High'])}")
        print(f"Total: {len(self.impact_stats['Medium']) + len(self.impact_stats['High'])}")
        print(f"{'='*60}\n")

    def print_intermediate_security_stats(self, file_name: str):
        """Print intermediate security statistics"""
        print(f"\n{'='*60}")
        print(f"Intermediate Security Statistics for: {file_name}")
        print(f"{'='*60}")
        print(f"Highest Issues: {len(self.security_stats['Highest'])}")
        print(f"High Issues: {len(self.security_stats['High'])}")
        print(f"Total: {len(self.security_stats['Highest']) + len(self.security_stats['High'])}")
        print(f"{'='*60}\n")

    def print_intermediate_triage_stats(self, file_name: str):
        """Print intermediate triage statistics"""
        print(f"\n{'='*60}")
        print(f"Intermediate Triage Statistics for: {file_name}")
        print(f"{'='*60}")
        print(f"Triaged Issues (without comment): {len(self.triage_stats['Triaged'])}")
        print(f"Dismissed Issues (without comment): {len(self.triage_stats['Dismissed'])}")
        print(f"Total: {len(self.triage_stats['Triaged']) + len(self.triage_stats['Dismissed'])}")
        print(f"{'='*60}\n")

    def generate_coverity_link(self, cids: Set[int]) -> str:
        """Generate Coverity link or message based on number of CIDs"""
        if len(cids) == 0:
            return "No issues"
        elif len(cids) <= 50:
            sorted_cids = sorted(cids)
            cid_params = "&".join([f"cid={cid}" for cid in sorted_cids])
            return f"[View in Coverity](https://{self.coverity_server}/query/defects.htm?project={self.project}&{cid_params})"
        else:
            return "Too many issues to display link"

    def generate_coverity_link_console(self, cids: Set[int]) -> str:
        """Generate Coverity link for console output (plain URL)"""
        if len(cids) == 0:
            return "No issues"
        elif len(cids) <= 50:
            sorted_cids = sorted(cids)
            cid_params = "&".join([f"cid={cid}" for cid in sorted_cids])
            return f"https://{self.coverity_server}/query/defects.htm?project={self.project}&{cid_params}"
        else:
            return "Too many issues to display link"

    def print_final_impact_stats(self):
        """Print final impact statistics to console"""
        print("\n" + "="*80)
        print("FINAL IMPACT-BASED STATISTICS")
        print("="*80)

        total = len(self.impact_stats['Medium']) + len(self.impact_stats['High'])

        print(f"\n{'Impact Level':<15} {'Count':<10} {'Coverity Link'}")
        print("-" * 80)

        for impact in ['High', 'Medium']:
            count = len(self.impact_stats[impact])
            link = self.generate_coverity_link_console(self.impact_stats[impact])
            print(f"{impact:<15} {count:<10} {link}")

        print("-" * 80)
        print(f"{'Total':<15} {total:<10}")
        print("="*80 + "\n")

    def print_final_security_stats(self):
        """Print final security statistics to console"""
        print("\n" + "="*80)
        print("FINAL SECURITY CLASSIFICATION STATISTICS")
        print("="*80)

        total = len(self.security_stats['Highest']) + len(self.security_stats['High'])

        print(f"\n{'Classification':<20} {'Count':<10} {'Coverity Link'}")
        print("-" * 80)

        for classification in ['Highest', 'High']:
            count = len(self.security_stats[classification])
            level = 1 if classification == 'Highest' else 2
            label = f"{classification} (Level {level})"
            link = self.generate_coverity_link_console(self.security_stats[classification])
            print(f"{label:<20} {count:<10} {link}")

        print("-" * 80)
        print(f"{'Total':<20} {total:<10}")
        print("="*80 + "\n")

    def print_final_triage_stats(self):
        """Print final triage statistics to console"""
        print("\n" + "="*80)
        print("FINAL TRIAGE STATISTICS (without comment)")
        print("="*80)

        total = len(self.triage_stats['Triaged']) + len(self.triage_stats['Dismissed'])

        print(f"\n{'Status':<15} {'Count':<10} {'Coverity Link'}")
        print("-" * 80)

        for status in ['Triaged', 'Dismissed']:
            count = len(self.triage_stats[status])
            link = self.generate_coverity_link_console(self.triage_stats[status])
            print(f"{status:<15} {count:<10} {link}")

        print("-" * 80)
        print(f"{'Total':<15} {total:<10}")
        print("="*80 + "\n")

    def check_exit_status(self) -> int:
        """Check if script should exit with error based on findings"""
        exit_code = 0

        # Always check impact-based issues
        total_impact = len(self.impact_stats['Medium']) + len(self.impact_stats['High'])
        if total_impact > 0:
            self.log_error(f"Found {total_impact} impact-based issues (High: {len(self.impact_stats['High'])}, Medium: {len(self.impact_stats['Medium'])})")
            exit_code = 1

        # Check security issues if flag is set
        if self.fail_on_security:
            total_security = len(self.security_stats['Highest']) + len(self.security_stats['High'])
            if total_security > 0:
                self.log_error(f"Found {total_security} security classification issues (Highest: {len(self.security_stats['Highest'])}, High: {len(self.security_stats['High'])})")
                exit_code = 1

        # Check triage issues if flag is set
        if self.fail_on_triage:
            total_triage = len(self.triage_stats['Triaged']) + len(self.triage_stats['Dismissed'])
            if total_triage > 0:
                self.log_error(f"Found {total_triage} triage issues without comment (Triaged: {len(self.triage_stats['Triaged'])}, Dismissed: {len(self.triage_stats['Dismissed'])})")
                exit_code = 1

        return exit_code

    def auto_fit_columns(self, worksheet, df):
        """Auto-fit column widths based on column headers and content"""
        for idx, column in enumerate(df.columns, 1):
            column_letter = get_column_letter(idx)
            # Set width based on column header length
            header_length = len(str(column))

            # Special handling for 'cid' column - fit to largest number
            if column.lower() == 'cid':
                max_length = header_length
                for cell in worksheet[column_letter]:
                    try:
                        if cell.value is not None and cell.row > 1:  # Skip header
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                    except:
                        pass
                worksheet.column_dimensions[column_letter].width = max_length + 2
            else:
                # Add some padding for other columns
                worksheet.column_dimensions[column_letter].width = header_length + 2

    def format_cid_url_as_link(self, worksheet, df):
        """Format cid_url column as clickable hyperlinks"""
        # Find the cid_url column index
        if 'cid_url' not in df.columns:
            return

        cid_url_col_idx = df.columns.get_loc('cid_url') + 1
        cid_url_col_letter = get_column_letter(cid_url_col_idx)

        # Format each cell in the cid_url column as a hyperlink
        for row_idx in range(2, len(df) + 2):  # Start from row 2 (skip header)
            cell = worksheet[f'{cid_url_col_letter}{row_idx}']
            url = cell.value
            if url and isinstance(url, str) and url.startswith('http'):
                # Set the cell as a hyperlink
                cell.hyperlink = url
                cell.style = 'Hyperlink'
                # Optionally set font color to blue and underline
                cell.font = Font(color="0563C1", underline="single")

    def process_all_files(self):
        """Process all Excel files"""
        for excel_file in self.excel_files:
            df = self.read_excel(excel_file)

            # Process impact-based filtering
            df_impact = self.filter_impact_based(df.copy())
            self.process_impact_stats(df_impact, excel_file)

            # Process security-based filtering
            df_security = self.filter_security_based(df.copy())
            self.process_security_stats(df_security, excel_file)

            # Process triage-based filtering
            df_triage = self.filter_triage_based(df.copy())
            self.process_triage_stats(df_triage, excel_file)

    def generate_impact_report(self, output_file: str = "impact_statistics.md"):
        """Generate markdown report for impact-based statistics"""
        self.log_info(f"Generating impact report: {output_file}")

        with open(output_file, 'w') as f:
            f.write("# Coverity Impact-Based Statistics Report\n\n")

            f.write("## Statistics by Impact Level\n\n")
            f.write("| Impact Level | Count | Coverity Link |\n")
            f.write("|--------------|-------|---------------|\n")

            total = len(self.impact_stats['Medium']) + len(self.impact_stats['High'])

            for impact in ['High', 'Medium']:
                count = len(self.impact_stats[impact])
                link = self.generate_coverity_link(self.impact_stats[impact])
                f.write(f"| {impact} | {count} | {link} |\n")

            f.write(f"| **Total** | **{total}** | - |\n")

        self.log_info(f"Impact report generated: {output_file}")

    def generate_security_report(self, output_file: str = "security_statistics.md"):
        """Generate markdown report for security classification-based statistics"""
        self.log_info(f"Generating security report: {output_file}")

        with open(output_file, 'w') as f:
            f.write("# Coverity Security Classification Statistics Report\n\n")

            f.write("## Statistics by Security Classification\n\n")
            f.write("| Classification | Count | Coverity Link |\n")
            f.write("|----------------|-------|---------------|\n")

            total = len(self.security_stats['Highest']) + len(self.security_stats['High'])

            for classification in ['Highest', 'High']:
                count = len(self.security_stats[classification])
                level = 1 if classification == 'Highest' else 2
                link = self.generate_coverity_link(self.security_stats[classification])
                f.write(f"| {classification} (Level {level}) | {count} | {link} |\n")

            f.write(f"| **Total** | **{total}** | - |\n")

        self.log_info(f"Security report generated: {output_file}")

    def generate_triage_report(self, output_file: str = "triage_statistics.md"):
        """Generate markdown report for triage-based statistics"""
        self.log_info(f"Generating triage report: {output_file}")

        with open(output_file, 'w') as f:
            f.write("# Coverity Triage Statistics Report (Without Comment)\n\n")

            f.write("## Statistics by Status\n\n")
            f.write("| Status | Count | Coverity Link |\n")
            f.write("|--------|-------|---------------|\n")

            total = len(self.triage_stats['Triaged']) + len(self.triage_stats['Dismissed'])

            for status in ['Triaged', 'Dismissed']:
                count = len(self.triage_stats[status])
                link = self.generate_coverity_link(self.triage_stats[status])
                f.write(f"| {status} | {count} | {link} |\n")

            f.write(f"| **Total** | **{total}** | - |\n")

        self.log_info(f"Triage report generated: {output_file}")

    def generate_excel_report(self, output_file: str = "filtered_issues.xlsx"):
        """Generate Excel file with all filtered issues"""
        self.log_info(f"Generating Excel report: {output_file}")

        try:
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                # Impact-based issues sheet
                if self.impact_issues:
                    df_impact = pd.DataFrame(self.impact_issues)
                    # Remove duplicates based on CID
                    df_impact = df_impact.drop_duplicates(subset=['cid'])
                    # Sort by impact (High first) and then by CID
                    df_impact['impact_order'] = df_impact['displayImpact'].map({'High': 0, 'Medium': 1})
                    df_impact = df_impact.sort_values(['impact_order', 'cid'])
                    df_impact = df_impact.drop('impact_order', axis=1)
                    df_impact.to_excel(writer, sheet_name='Impact_Issues', index=False)

                    # Get the worksheet and apply formatting
                    worksheet = writer.sheets['Impact_Issues']
                    self.auto_fit_columns(worksheet, df_impact)
                    self.format_cid_url_as_link(worksheet, df_impact)

                    self.log_info(f"Added {len(df_impact)} impact-based issues to Excel")
                else:
                    # Create empty sheet
                    pd.DataFrame().to_excel(writer, sheet_name='Impact_Issues', index=False)
                    self.log_info("No impact-based issues found")

                # Security classification issues sheet
                if self.security_issues:
                    df_security = pd.DataFrame(self.security_issues)
                    # Remove duplicates based on CID
                    df_security = df_security.drop_duplicates(subset=['cid'])
                    # Sort by classification (Highest first) and then by CID
                    df_security['security_order'] = df_security['security_classification'].map({'1 - Highest': 0, '2 - High': 1})
                    df_security = df_security.sort_values(['security_order', 'cid'])
                    df_security = df_security.drop('security_order', axis=1)
                    df_security.to_excel(writer, sheet_name='Security_Issues', index=False)

                    # Get the worksheet and apply formatting
                    worksheet = writer.sheets['Security_Issues']
                    self.auto_fit_columns(worksheet, df_security)
                    self.format_cid_url_as_link(worksheet, df_security)

                    self.log_info(f"Added {len(df_security)} security classification issues to Excel")
                else:
                    # Create empty sheet
                    pd.DataFrame().to_excel(writer, sheet_name='Security_Issues', index=False)
                    self.log_info("No security classification issues found")

                # Triage issues sheet (without comment)
                if self.triage_issues:
                    df_triage = pd.DataFrame(self.triage_issues)
                    # Remove duplicates based on CID
                    df_triage = df_triage.drop_duplicates(subset=['cid'])
                    # Sort by status (Triaged first) and then by CID
                    df_triage['status_order'] = df_triage['status'].map({'Triaged': 0, 'Dismissed': 1})
                    df_triage = df_triage.sort_values(['status_order', 'cid'])
                    df_triage = df_triage.drop('status_order', axis=1)
                    df_triage.to_excel(writer, sheet_name='Triage_Issues', index=False)

                    # Get the worksheet and apply formatting
                    worksheet = writer.sheets['Triage_Issues']
                    self.auto_fit_columns(worksheet, df_triage)
                    self.format_cid_url_as_link(worksheet, df_triage)

                    self.log_info(f"Added {len(df_triage)} triage issues (without comment) to Excel")
                else:
                    # Create empty sheet
                    pd.DataFrame().to_excel(writer, sheet_name='Triage_Issues', index=False)
                    self.log_info("No triage issues (without comment) found")

            self.log_info(f"Excel report generated: {output_file}")
        except Exception as e:
            self.log_error(f"Failed to generate Excel report: {e}")

def main():
    parser = argparse.ArgumentParser(
        description="Parse Coverity Excel files and generate statistics for MPP project"
    )
    parser.add_argument(
        "files",
        nargs="+",
        help="Excel file(s) to parse"
    )
    parser.add_argument(
        "-d", "--debug",
        action="store_true",
        help="Enable debug logging (prints full row data)"
    )
    parser.add_argument(
        "-i", "--intermediate",
        action="store_true",
        help="Show intermediate statistics for each file"
    )
    parser.add_argument(
        "-m", "--impact-output",
        default=None,
        help="Output file for impact statistics (default: coverity_output/impact_statistics.md)"
    )
    parser.add_argument(
        "-s", "--security-output",
        default=None,
        help="Output file for security statistics (default: coverity_output/security_statistics.md)"
    )
    parser.add_argument(
        "-t", "--triage-output",
        default=None,
        help="Output file for triage statistics (default: coverity_output/triage_statistics.md)"
    )
    parser.add_argument(
        "-e", "--excel-output",
        default=None,
        help="Output Excel file for filtered issues (default: coverity_output/filtered_issues.xlsx)"
    )
    parser.add_argument(
        "-c", "--coverity-server",
        default="coverity9.nxp.com",
        help="Coverity server URL (default: coverity9.nxp.com)"
    )
    parser.add_argument(
        "-p", "--project",
        default="MPP",
        help="Coverity project name (default: MPP)"
    )
    parser.add_argument(
        "-f", "--fail-on-security",
        action="store_true",
        help="Exit with error code if security classification issues (Highest/High) are found"
    )
    parser.add_argument(
        "-g", "--fail-on-triage",
        action="store_true",
        help="Exit with error code if triage issues without comment (Triaged/Dismissed) are found"
    )

    args = parser.parse_args()

    # Validate files exist
    for file_path in args.files:
        if not Path(file_path).exists():
            print(f"[ERROR] File not found: {file_path}")
            sys.exit(1)

    # Determine output directory
    output_dir = None
    provided_outputs = [args.impact_output, args.security_output, args.triage_output, args.excel_output]

    # Check if any output path was provided
    for output_path in provided_outputs:
        if output_path is not None:
            path_obj = Path(output_path)
            # If it's a directory, use it as output_dir
            if path_obj.is_dir():
                output_dir = path_obj
                break
            # If it's a file path, use its parent directory
            elif path_obj.parent != Path('.'):
                output_dir = path_obj.parent
                break

    # If no output directory was determined, use default
    if output_dir is None:
        output_dir = Path("coverity_output")

    # Create output directory if it doesn't exist
    output_dir.mkdir(parents=True, exist_ok=True)

    # Helper function to resolve output path
    def resolve_output_path(provided_path, default_filename):
        if provided_path is None:
            return output_dir / default_filename

        path_obj = Path(provided_path)

        # If it's a directory, append default filename
        if path_obj.is_dir():
            return path_obj / default_filename

        # If parent is current directory and no directory was specified, use output_dir
        if path_obj.parent == Path('.') and output_dir != Path('.'):
            return output_dir / path_obj.name

        # Otherwise use the provided path as-is
        return path_obj

    # Resolve all output paths
    impact_output = resolve_output_path(args.impact_output, "impact_statistics.md")
    security_output = resolve_output_path(args.security_output, "security_statistics.md")
    triage_output = resolve_output_path(args.triage_output, "triage_statistics.md")
    excel_output = resolve_output_path(args.excel_output, "filtered_issues.xlsx")

    # Create parser and process files
    coverity_parser = CoverityParser(
        excel_files=args.files,
        debug=args.debug,
        show_intermediate=args.intermediate,
        coverity_server=args.coverity_server,
        project=args.project,
        fail_on_security=args.fail_on_security,
        fail_on_triage=args.fail_on_triage
    )

    coverity_parser.process_all_files()

    # Print final statistics to console
    coverity_parser.print_final_impact_stats()
    coverity_parser.print_final_security_stats()
    coverity_parser.print_final_triage_stats()

    # Generate reports
    coverity_parser.generate_impact_report(str(impact_output))
    coverity_parser.generate_security_report(str(security_output))
    coverity_parser.generate_triage_report(str(triage_output))
    coverity_parser.generate_excel_report(str(excel_output))

    # Check exit status
    exit_code = coverity_parser.check_exit_status()

    print("\n" + "="*60)
    if exit_code == 0:
        print("Processing Complete! No issues found.")
    else:
        print("Processing Complete! Issues found - exiting with error.")
    print("="*60)
    print(f"Impact report: {impact_output}")
    print(f"Security report: {security_output}")
    print(f"Triage report: {triage_output}")
    print(f"Excel report: {excel_output}")
    print("="*60)

    sys.exit(exit_code)

if __name__ == "__main__":
    main()
